"""The made-up old Time-Off Sheet (fixtures/old-sheet.json) as Google Sheets downloads it: an .xlsx whose dates are real
date cells and timestamps real date-times, numbers numbers — and one CSV per tab ("NLO-Time off - <Tab>.csv").
Made-up people only. Run after make_fixture.js."""
import csv, datetime, json, pathlib, re
import openpyxl
here = pathlib.Path(__file__).parent / 'fixtures'
data = json.loads((here / 'old-sheet.json').read_text())
D = re.compile(r'^\d{4}-\d{2}-\d{2}$'); DT = re.compile(r'^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$')
def cell(v):
    if isinstance(v, str) and D.match(v): return datetime.date.fromisoformat(v)
    if isinstance(v, str) and DT.match(v): return datetime.datetime.fromisoformat(v)
    return v
wb = openpyxl.Workbook(); wb.remove(wb.active)
for name, rows in data['tabs'].items():
    ws = wb.create_sheet(name)
    for r in rows: ws.append([cell(v) for v in r])
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, datetime.datetime): c.number_format = 'yyyy-mm-dd hh:mm:ss'
            elif isinstance(c.value, datetime.date): c.number_format = 'yyyy-mm-dd'
    with open(here / ('NLO-Time off - ' + name + '.csv'), 'w', newline='') as f:
        w = csv.writer(f)
        for r in rows: w.writerow(['' if v is None else v for v in r])
wb.save(here / 'old-sheet.xlsx')
print('wrote old-sheet.xlsx and', len(data['tabs']), 'CSV files')
