"""Made-up Edge "Accounts Receivable Aging" exports for the tests: the same layout as a real export
(title + run time, date line, options, subgroup, headings at the same columns, merged cells, "Days" as text,
the billing note past the last column, the "(N Patients)" total line, the "Edge v…" line), as .xls (BIFF8),
.xlsx and tab-separated text (what Excel puts on the clipboard). Every name and number is invented.
Run: python3 test/make_fixtures.py  → test/fixtures/*  (needs xlwt and openpyxl)"""
import datetime, json, os, random
import xlwt, openpyxl
from openpyxl.styles import numbers

OUT = os.path.join(os.path.dirname(__file__), 'fixtures')
os.makedirs(OUT, exist_ok=True)
random.seed(20261005)
FIRST = ['Avery', 'Riley', 'Quinn', 'Harper', 'Rowan', 'Skyler', 'Finley', 'Parker', 'Emerson', 'Sawyer', 'Jordan', 'Dana', 'Reese', 'Logan',
         'Blake', 'Peyton', 'Drew', 'Kendall', 'Hayden', 'Dakota', 'Elliot', 'Cameron', 'Ellis', 'Marlow', 'Shay', 'Arden', 'Remy', 'Sage', 'Zoë', 'José', 'Łucja']
LAST = ['Sample', 'Demo', 'Example', 'Testcase', 'Placeholder', 'Mockley', 'Fakeworth', 'Exampleton', 'Samplesen', 'Demoray', 'Notreal', 'Exämple', "O'Sample", 'Demo-Ray']
STS = ['A-Comp Bra', 'A-Comp Bra', 'A-Comp Bra', 'A-Comp Inv', 'A-Comp NLA', 'A- Ph I Br', 'A- Ph II B', 'Retention', 'Inactive', 'RET ONLY', 'Start Sche']
NOTES = ['Payment in process', 'Autopay card expired', 'Sent letter 9/14', 'Plan — $50/mo']
AS_OF = datetime.date(2026, 10, 5)
RUN = datetime.datetime(2026, 10, 5, 14, 8, 22)
HEAD = [(0, 'Patient'), (1, 'Sts'), (2, 'Responsible Party'), (4, 'Home  Ph'), (5, 'Work  Ph'), (6, 'Amt Due'), (7, '0-30'), (8, '31-60'),
        (9, '61-90'), (10, '91+'), (12, 'Days'), (14, 'Balance'), (16, 'Last Amt'), (17, 'Received')]

def r2(x): return round(x + 1e-9, 2)
used = set()
def name():
    while True:
        n = random.choice(FIRST) + ' ' + random.choice(LAST)
        if n not in used: used.add(n); return n

def account(i):
    pat = name(); last = pat.split(' ', 1)[1]
    ins = i % 5 == 3
    rp = ('INS: ' if ins else '') + random.choice(['', '', 'Mr. ', 'Mrs. ']) + random.choice(['Jordan', 'Alex', 'Pat', 'Chris', 'Robin', 'Lee']) + ' ' + last
    home = '(352) 555-01%02d' % (i % 100)
    work = '(352) 555-01%02d' % ((i * 7) % 100) if i % 23 == 0 else ''
    kind = random.random()
    if kind < 0.18:  # credit balance
        credit = r2(random.choice([25, 62.5, 85, 128, 240, 310, 450, 1150]) + random.choice([0, 0.25, 0.5]))
        due = -credit; bk = [0, 0, 0, 0]; days = ''; bal = -credit
        age = random.choice([12, 40, 95, 220, 410, 800, 1290])
        last_amt = r2(credit + random.choice([0, 100, 150])); recv = AS_OF - datetime.timedelta(days=age)
    else:
        m = 11.11 if ins else random.choice([45, 150, 175, 199, 225, 250, 275, 300])
        d = random.choice([4, 9, 15, 24, 33, 47, 58, 66, 79, 93, 112, 141, 186, 233, 301, 395, 488])
        n = -(-d // 30)
        bk = [m, m if n >= 2 else 0, m if n >= 3 else 0, (n - 3) * m if n > 3 else 0]
        if ins and i % 10 == 8: bk[3 if d > 90 else 0] += 125  # an unpaid claim on top: partly paid
        bk = [r2(x) for x in bk]
        due = r2(sum(bk)); days = str(d); bal = r2(due + m * random.randint(3, 20))
        never = ins and i % 10 != 8 and i % 3 == 0
        last_amt = None if never else m
        recv = None if never else AS_OF - datetime.timedelta(days=d + 3)
    note = random.choice(NOTES) if i % 11 == 4 else ''
    return dict(patient=pat, sts=random.choice(STS), rp=rp, home=home, work=work, due=due, b0=bk[0], b30=bk[1], b60=bk[2], b90=bk[3],
                days=days, bal=bal, lastAmt=last_amt, recv=recv, note=note)

def build(n_rows, subgroup, options, repeat_header_at=None):
    accts = [account(i) for i in range(n_rows)]
    rows = []  # (row index → list of (col, value, kind)) where kind: s string, n number, d date, dt datetime
    rows.append([(0, 'Accounts Receivable Aging', 's'), (13, RUN, 'dt')])
    rows.append([(0, 'Accounts Receivable Aging', 's')])
    rows.append([(0, 'Monday, October 5, 2026  Office: All,  Doctor: All', 's')])
    rows.append([(0, options, 's')])
    rows.append([(0, 'Subgroup: ' + subgroup, 's')])
    rows.append([]); rows.append([])
    head = [(c, t, 's') for c, t in HEAD]
    rows.append(head)
    for i, a in enumerate(accts):
        if repeat_header_at and i == repeat_header_at:
            rows.append([(0, 'Accounts Receivable Aging', 's')]); rows.append(head)
        r = [(0, a['patient'], 's'), (1, a['sts'], 's'), (2, a['rp'], 's'), (4, a['home'], 's')]
        if a['work']: r.append((5, a['work'], 's'))
        r += [(6, a['due'], 'n'), (7, a['b0'], 'n'), (8, a['b30'], 'n'), (9, a['b60'], 'n'), (10, a['b90'], 'n')]
        if a['days'] != '': r.append((12, a['days'], 's'))
        r.append((14, a['bal'], 'n'))
        if a['lastAmt'] is not None: r.append((16, a['lastAmt'], 'n'))
        if a['recv'] is not None: r.append((17, a['recv'], 'd'))
        if a['note']: r.append((19, a['note'], 's'))
        rows.append(r)
    tot = {k: r2(sum(a[k] for a in accts)) for k in ['due', 'b0', 'b30', 'b60', 'b90', 'bal']}
    rows.append([(0, '(%d Patients)' % len(accts), 's'), (6, tot['due'], 'n'), (7, tot['b0'], 'n'), (8, tot['b30'], 'n'), (9, tot['b60'], 'n'), (10, tot['b90'], 'n'), (14, tot['bal'], 'n')])
    rows.append([(0, 'Edge v7.2.3.20006', 's')])
    expected = [dict(patient=a['patient'], sts=a['sts'], rp=a['rp'], home=a['home'], work=a['work'], due=a['due'], b0=a['b0'], b30=a['b30'], b60=a['b60'], b90=a['b90'],
                     days=int(a['days']) if a['days'] != '' else None, bal=a['bal'], lastAmt=a['lastAmt'], recv=a['recv'].isoformat() if a['recv'] else '', note=a['note']) for a in accts]
    return rows, expected, tot

def merges_for(r_index, row):
    # Edge's merged cells: the party name over two columns, amounts and days over two
    out = []
    cols = {c for c, _, _ in row}
    for c0, c1 in [(2, 3), (10, 11), (12, 13), (14, 15)]:
        if c0 in cols and r_index % 3 == 0: out.append((r_index, r_index, c0, c1))
    return out

def write_xls(path, rows):
    wb = xlwt.Workbook(encoding='utf-8')
    ws = wb.add_sheet('Sheet1'); wb.add_sheet('Sheet2'); wb.add_sheet('Sheet3')
    money = xlwt.easyxf(num_format_str='#,##0.00'); date = xlwt.easyxf(num_format_str='M/D/YYYY'); dt = xlwt.easyxf(num_format_str='M/D/YYYY h:mm AM/PM')
    for ri, row in enumerate(rows):
        merged = merges_for(ri, row) if ri > 7 else ([(3, 3, 0, 19)] if ri == 3 else [])
        mcols = {m[2] for m in merged}
        for c, v, k in row:
            st = money if k == 'n' else date if k == 'd' else dt if k == 'dt' else xlwt.Style.default_style
            if c in mcols:
                m = [x for x in merged if x[2] == c][0]; ws.write_merge(m[0], m[1], m[2], m[3], v, st)
            else: ws.write(ri, c, v, st)
    wb.save(path)

def write_xlsx(path, rows):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Sheet1'; wb.create_sheet('Sheet2')
    for ri, row in enumerate(rows):
        for c, v, k in row:
            cell = ws.cell(row=ri + 1, column=c + 1, value=v)
            if k == 'n': cell.number_format = '#,##0.00'
            if k == 'd': cell.number_format = 'm/d/yyyy'
            if k == 'dt': cell.number_format = 'm/d/yyyy h:mm AM/PM'
        if ri > 7:
            for m in merges_for(ri, row): ws.merge_cells(start_row=m[0] + 1, end_row=m[1] + 1, start_column=m[2] + 1, end_column=m[3] + 1)
    wb.save(path)

def write_tsv(path, rows):
    # what Excel puts on the clipboard: tab-separated, dates and amounts as shown on screen
    lines = []
    for row in rows:
        width = max([c for c, _, _ in row] + [0]) + 1; cells = [''] * width
        for c, v, k in row:
            if k == 'n': cells[c] = '{:,.2f}'.format(v) if v >= 0 else '({:,.2f})'.format(-v)
            elif k == 'd': cells[c] = '%d/%d/%d' % (v.month, v.day, v.year)
            elif k == 'dt': cells[c] = '%d/%d/%d %d:%02d PM' % (v.month, v.day, v.year, v.hour - 12, v.minute)
            else: cells[c] = str(v)
        lines.append('\t'.join(cells))
    with open(path, 'w', encoding='utf-8') as f: f.write('\r\n'.join(lines) + '\r\n')

def write_html(path, rows):
    # what Excel puts on the clipboard as HTML: merged cells as colspan, amounts as shown, &nbsp; for blanks
    from html import escape
    out = ['<html><body><table>']
    for ri, row in enumerate(rows):
        width = max([c for c, _, _ in row] + [0]) + 1; cells = {c: (v, k) for c, v, k in row}
        spans = {m[2]: m[3] - m[2] + 1 for m in merges_for(ri, row)} if ri > 7 else {}
        tds, c = [], 0
        while c < max(width, 20):
            if c in cells:
                v, k = cells[c]
                t = ('{:,.2f}'.format(v) if v >= 0 else '({:,.2f})'.format(-v)) if k == 'n' else ('%d/%d/%d' % (v.month, v.day, v.year)) if k == 'd' else ('%d/%d/%d %d:%02d PM' % (v.month, v.day, v.year, v.hour - 12, v.minute)) if k == 'dt' else escape(str(v))
                sp = spans.get(c, 1)
                tds.append('<td' + (' colspan=%d' % sp if sp > 1 else '') + ' class=xl65>' + t + '</td>'); c += sp
            else: tds.append('<td>&nbsp;</td>'); c += 1
        out.append('<tr>' + ''.join(tds) + '</tr>')
    out.append('</table></body></html>')
    with open(path, 'w', encoding='utf-8') as f: f.write('\n'.join(out))

full_rows, full_exp, full_tot = build(320, 'None', 'Exclude Zero Dollar Balances', repeat_header_at=170)
write_html(os.path.join(OUT, 'edge-full.html'), full_rows)
write_xls(os.path.join(OUT, 'edge-full.xls'), full_rows)
write_xlsx(os.path.join(OUT, 'edge-full.xlsx'), full_rows)
write_tsv(os.path.join(OUT, 'edge-full.tsv'), full_rows)
random.seed(7); used.clear()
pd_rows, pd_exp, pd_tot = build(40, 'Past Due', 'Exclude Zero Dollar Balances, Exclude Insurance Contracts')
write_xls(os.path.join(OUT, 'edge-pastdue-noins.xls'), pd_rows)
with open(os.path.join(OUT, 'expected.json'), 'w') as f:
    json.dump({'full': {'rows': full_exp, 'total': full_tot, 'asOf': AS_OF.isoformat()}, 'pastdue': {'rows': pd_exp, 'total': pd_tot}}, f, indent=0, ensure_ascii=False)
print('wrote', len(full_exp), 'and', len(pd_exp), 'made-up accounts')
